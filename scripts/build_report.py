"""Assemble the expanded Fedfina report. Usage: python3 build_report.py [--parts 1,2,3] [--out path]"""
import os, sys, argparse, importlib
sys.path.insert(0, os.path.dirname(__file__))
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from report_builder import new_document, render, Page, NAVY, TEAL, PAGE_HEIGHT_IN
import report_charts

CHARTS = "/tmp/work/charts"
DEFAULT_OUT = "/home/user/Datta/Summer Internship Project final By DP (expanded with 5-year data).docx"


def cover(doc):
    for _ in range(6):
        doc.add_paragraph()
    def line(text, size, color=None, bold=True, after=6):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text); r.bold = bold; r.font.size = Pt(size)
        if color: r.font.color.rgb = color
        p.paragraph_format.space_after = Pt(after)
        return p
    line("A Project Report on", 12, bold=False, after=10)
    line("FUNDAMENTAL ANALYSIS OF", 24, NAVY, after=0)
    line("FEDBANK FINANCIAL SERVICES LIMITED", 24, NAVY, after=12)
    line("A five-year study (FY 2021-22 to FY 2025-26) of a bank-promoted secured retail NBFC, prepared from the company’s Annual Reports FY24, FY25 and FY26", 12, TEAL, bold=False, after=24)
    line("Submitted in partial fulfilment of the requirements for the", 11, bold=False, after=0)
    line("Master of Management Studies (MMS) – Finance", 13, after=0)
    line("University of Mumbai", 11, bold=False, after=24)
    line("Submitted by", 11, bold=False, after=0)
    line("Dattaguru Kallapa Patil", 14, after=0)
    line("Roll No. M-16144  |  Batch 2026–27", 11, bold=False, after=24)
    line("Under the guidance of", 11, bold=False, after=0)
    line("Dr. Rahul N. Wadekar  and  Dr. Prasad Supekar", 13, after=24)
    line("DES’s Navinchandra Mehta Institute of Technology and Development", 12, NAVY, after=0)
    line("Dadar (W), Mumbai – 400028", 11, bold=False, after=0)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--parts", default="1,2,3")
    ap.add_argument("--out", default=DEFAULT_OUT)
    ap.add_argument("--no-charts", action="store_true")
    a = ap.parse_args()
    if not a.no_charts:
        report_charts.make_all(CHARTS)
    pages = []
    for n in a.parts.split(","):
        mod = importlib.import_module(f"content_part{n.strip()}")
        pages += mod.pages(CHARTS)
    doc = new_document()
    cover(doc)
    warn = 0
    for i, pg in enumerate(pages):
        render(doc, pg)
        flag = ""
        if pg.height > 9.3: flag = "  <-- OVERFLOW RISK"; warn += 1
        elif pg.height < 7.6: flag = "  <-- UNDERFILLED"; warn += 1
        print(f"{i+2:3d}  {pg.height:5.2f}in  {pg.title[:60]}{flag}")
    doc.save(a.out)
    print(f"\nSections: {len(pages)} (+cover). Flags: {warn}. Saved: {a.out}")


if __name__ == "__main__":
    main()
