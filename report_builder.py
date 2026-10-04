#!/usr/bin/env python3
"""Build the academic report and companion risk-monitoring workbook.

Inputs are the three supplied Fedfina annual reports and Rashmi Gupta's
reference project PDF. The employer-certificate page is intentionally not
forged: it is a student-prepared internship-particulars page.
"""
from __future__ import annotations

import html
import math
import os
import tempfile
from pathlib import Path

import pymupdf
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np

from report_pages import PAGES

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas
from reportlab.platypus import Image as RLImage, Paragraph, Table, TableStyle
from reportlab.lib.utils import ImageReader

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from openpyxl import Workbook, load_workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent
REFERENCE_PDF = ROOT / "Rashmi Gupta Project.pdf"
FY24_PDF = ROOT / "Annual reports of Fedbank FY 24.pdf"
FY25_PDF = ROOT / "Annual reports of Fedbank FY 25.pdf"
FY26_PDF = ROOT / "Annual reports of Fedbank FY 26.pdf"
DOCX_OUT = ROOT / "Summer Intership Project final By DP.docx"
PDF_OUT = ROOT / "Risk Analysis at Fedbank Financial Services Ltd - Dattaguru Patil.pdf"
XLSX_OUT = ROOT / "Fedfina Risk Analyst Internship Workbook.xlsx"

NAVY = "#17466B"
BLUE = "#2F759D"
ORANGE = "#E8A23A"
PALE = "#EAF2F7"
INK = "#202B35"
MUTED = "#66737E"
GRID = "#D5DEE5"
RED = "#B84A4A"
GREEN = "#4C8969"

FYS = ["FY22", "FY23", "FY24", "FY25", "FY26"]
AUM_CR = [6187.2, 9069.6, 12191.9, 15811.5, 20153.0]
PAT_CR = [103.5, 180.1, 244.7, 225.2, 343.6]
YIELD = [16.0, 16.1, 16.7, 17.1, 16.7]
SPREAD = [8.2, 8.3, 8.1, 8.2, 8.6]
COF = [7.8, 7.8, 8.6, 9.0, 8.1]
COST_INCOME = [58.4, 58.6, 58.2, 57.6, 57.2]
ROA = [1.7, 2.3, 2.4, 1.8, 2.4]
ROE = [10.4, 14.4, 13.5, 9.4, 12.6]
GNPA = [2.2, 2.0, 1.7, 2.0, 1.9]
NNPA = [1.8, 1.6, 1.3, 1.2, 1.3]
PCR = [22.1, 22.2, 20.4, 40.0, 32.3]
CRAR = [23.0, 17.9, 23.5, 21.9, 22.4]


def extract_reference_assets(workdir: Path) -> tuple[Path, Path]:
    """Extract the institutional crest and letterhead used by the reference."""
    doc = pymupdf.open(REFERENCE_PDF)
    found = []
    for page_index in (0, 1):
        page = doc[page_index]
        images = page.get_images(full=True)
        if not images:
            raise RuntimeError(f"No reference image found on page {page_index + 1}")
        xref = images[0][0]
        payload = doc.extract_image(xref)
        out = workdir / f"reference_page_{page_index+1}.{payload['ext']}"
        out.write_bytes(payload["image"])
        found.append(out)
    return found[0], found[1]


def set_plot_style(ax):
    ax.grid(axis="y", color="#E5EAF0", linewidth=0.75)
    ax.set_axisbelow(True)
    for side in ("top", "right", "left"):
        ax.spines[side].set_visible(False)
    ax.spines["bottom"].set_color("#AAB5BF")
    ax.tick_params(axis="y", colors="#56636D", labelsize=8, length=0)
    ax.tick_params(axis="x", colors="#384550", labelsize=8, length=0)


def save_charts(workdir: Path) -> dict[str, Path]:
    charts: dict[str, Path] = {}
    plt.rcParams.update({
        "font.family": "DejaVu Sans", "font.size": 8.5,
        "axes.titleweight": "bold", "axes.titlesize": 10,
        "figure.facecolor": "white", "axes.facecolor": "white",
    })

    def finish(fig, key):
        fig.tight_layout(pad=1.4)
        path = workdir / f"chart_{key}.png"
        fig.savefig(path, dpi=170, bbox_inches="tight", facecolor="white")
        plt.close(fig)
        charts[key] = path

    x = np.arange(len(FYS))
    fig, ax = plt.subplots(figsize=(7.0, 2.4))
    bars = ax.bar(x, AUM_CR, color=["#AFC4D2"] * 4 + [NAVY], width=.62)
    set_plot_style(ax)
    ax.set_title("Assets Under Management | ₹ crore", loc="left", color=NAVY)
    ax.set_xticks(x, FYS)
    ax.set_ylim(0, 22500)
    ax.set_ylabel("₹ crore", color=MUTED, fontsize=8)
    for b, v in zip(bars, AUM_CR):
        ax.text(b.get_x()+b.get_width()/2, v+420, f"{v:,.0f}", ha="center", va="bottom", fontsize=7.6, color=INK)
    finish(fig, "aum")

    fig, ax1 = plt.subplots(figsize=(7.0, 2.45))
    bars = ax1.bar(x, PAT_CR, color=["#C8D8E1"] * 4 + [NAVY], width=.56, label="PAT")
    set_plot_style(ax1)
    ax1.set_title("Profit After Tax and Return on Assets", loc="left", color=NAVY)
    ax1.set_xticks(x, FYS)
    ax1.set_ylabel("PAT (₹ crore)", color=MUTED, fontsize=8)
    ax1.set_ylim(0, 400)
    for b, v in zip(bars, PAT_CR):
        ax1.text(b.get_x()+b.get_width()/2, v+7, f"{v:.0f}", ha="center", fontsize=7.2, color=INK)
    ax2 = ax1.twinx()
    ax2.plot(x, ROA, color=ORANGE, marker="o", linewidth=2, markersize=4, label="ROA")
    ax2.set_ylim(0, 3.2)
    ax2.set_ylabel("ROA (%)", color="#9A651C", fontsize=8)
    ax2.tick_params(axis="y", colors="#9A651C", labelsize=8, length=0)
    ax2.spines["top"].set_visible(False)
    ax2.spines["right"].set_visible(False)
    for xx, v in zip(x, ROA):
        ax2.text(xx, v+.10, f"{v:.1f}%", ha="center", color="#8B5B1A", fontsize=7)
    ax1.legend([bars, ax2.lines[0]], ["PAT", "ROA"], loc="upper left", frameon=False, fontsize=7)
    finish(fig, "profitability")

    fig, ax = plt.subplots(figsize=(7.0, 2.5))
    set_plot_style(ax)
    ax.set_title("Pricing and Funding-Cost Indicators | %", loc="left", color=NAVY)
    ax.plot(x, YIELD, color=NAVY, marker="o", linewidth=2, label="Yield")
    ax.plot(x, SPREAD, color=ORANGE, marker="o", linewidth=2, label="Spread")
    ax.plot(x, COF, color=BLUE, marker="o", linewidth=2, label="Cost of borrowings")
    ax.set_xticks(x, FYS)
    ax.set_ylim(6.5, 18.5)
    ax.set_ylabel("Percent", color=MUTED, fontsize=8)
    ax.legend(loc="upper left", frameon=False, ncol=3, fontsize=7.2)
    finish(fig, "pricing")

    fig, (ax1, ax2) = plt.subplots(1, 2, figsize=(7.0, 2.55))
    for ax in (ax1, ax2):
        set_plot_style(ax)
        ax.set_xticks(x, FYS)
    ax1.set_title("NPA Ratios", loc="left", color=NAVY)
    ax1.plot(x, GNPA, color=NAVY, marker="o", linewidth=2, label="GNPA")
    ax1.plot(x, NNPA, color=ORANGE, marker="o", linewidth=2, label="NNPA")
    ax1.set_ylim(0, 2.8)
    ax1.set_ylabel("Percent", color=MUTED, fontsize=8)
    ax1.legend(frameon=False, fontsize=7, loc="lower left")
    ax2.set_title("Provision Coverage Ratio", loc="left", color=NAVY)
    bars = ax2.bar(x, PCR, color=["#CAD8E0"] * 4 + [BLUE], width=.58)
    ax2.set_ylim(0, 48)
    for b, v in zip(bars, PCR):
        ax2.text(b.get_x()+b.get_width()/2, v+1, f"{v:.1f}", ha="center", fontsize=6.8, color=INK)
    finish(fig, "asset_quality")

    fig, ax1 = plt.subplots(figsize=(7.0, 2.5))
    bars = ax1.bar(x, CRAR, color=["#C8D8E1"] * 4 + [NAVY], width=.55, label="CRAR")
    set_plot_style(ax1)
    ax1.set_title("Capital Adequacy and Return on Equity", loc="left", color=NAVY)
    ax1.set_xticks(x, FYS)
    ax1.set_ylim(0, 30)
    ax1.set_ylabel("CRAR (%)", color=MUTED, fontsize=8)
    for b, v in zip(bars, CRAR):
        ax1.text(b.get_x()+b.get_width()/2, v+.55, f"{v:.1f}%", ha="center", fontsize=7, color=INK)
    ax2 = ax1.twinx()
    ax2.plot(x, ROE, color=ORANGE, marker="o", linewidth=2, label="ROE")
    ax2.set_ylim(0, 18)
    ax2.set_ylabel("ROE (%)", color="#9A651C", fontsize=8)
    ax2.tick_params(axis="y", colors="#9A651C", labelsize=8, length=0)
    ax2.spines["top"].set_visible(False)
    ax2.spines["right"].set_visible(False)
    for xx, v in zip(x, ROE):
        ax2.text(xx, v+.55, f"{v:.1f}%", ha="center", fontsize=7, color="#8B5B1A")
    ax1.legend([bars, ax2.lines[0]], ["CRAR", "ROE"], loc="upper left", frameon=False, fontsize=7)
    finish(fig, "capital")

    fig, ax = plt.subplots(figsize=(7.0, 2.3))
    labels = ["FY24", "FY25", "FY26"]
    secure = [85.0, 89.5, 98.9]
    bars = ax.bar(labels, secure, color=["#AFC4D2", "#7EA8C1", NAVY], width=.52)
    set_plot_style(ax)
    ax.set_title("Secured AUM Share | As Reported", loc="left", color=NAVY)
    ax.set_ylim(0, 108)
    ax.set_ylabel("Percent of AUM", color=MUTED, fontsize=8)
    for b, v in zip(bars, secure):
        ax.text(b.get_x()+b.get_width()/2, v+1.3, f"{v:.1f}%", ha="center", fontsize=8, color=INK)
    finish(fig, "secured_mix")

    fig, ax = plt.subplots(figsize=(7.0, 2.15))
    lab = ["Within 1 year", "After 1 year", "Total"]
    vals = [4976.24, -2050.13, 2926.08]
    cols = [BLUE, RED, NAVY]
    ax.barh(np.arange(3), vals, color=cols, height=.5)
    ax.axvline(0, color="#697782", linewidth=.8)
    ax.set_yticks(np.arange(3), lab)
    ax.invert_yaxis()
    ax.set_xlim(-2800, 5800)
    ax.set_title("Net Position in Disclosed Maturity Analysis | ₹ crore", loc="left", color=NAVY)
    ax.grid(axis="x", color="#E5EAF0", linewidth=.75)
    ax.set_axisbelow(True)
    ax.spines["top"].set_visible(False); ax.spines["right"].set_visible(False); ax.spines["left"].set_visible(False)
    ax.tick_params(axis="x", labelsize=7.5, colors=MUTED, length=0)
    ax.tick_params(axis="y", labelsize=8, colors=INK, length=0)
    for yi, v in enumerate(vals):
        ax.text(v+(110 if v>=0 else -110), yi, f"{v:+,.0f}", va="center", ha="left" if v>=0 else "right", fontsize=8, color=INK)
    finish(fig, "liquidity")

    fig, ax = plt.subplots(figsize=(7.0, 2.2))
    labels = ["Loans", "Borrowings"]
    values = [6.52, -14.79]
    bars = ax.barh(labels, values, color=[BLUE, RED], height=.45)
    ax.axvline(0, color="#697782", linewidth=.8)
    ax.set_xlim(-18, 10)
    ax.set_title("Disclosed PAT Sensitivity to +25 bp Shock | ₹ crore", loc="left", color=NAVY)
    ax.grid(axis="x", color="#E5EAF0", linewidth=.75)
    ax.set_axisbelow(True)
    for side in ("top", "right", "left"):
        ax.spines[side].set_visible(False)
    ax.tick_params(axis="x", colors=MUTED, labelsize=8, length=0)
    ax.tick_params(axis="y", colors=INK, labelsize=8, length=0)
    for b, v in zip(bars, values):
        ax.text(v+(0.4 if v>=0 else -0.4), b.get_y()+b.get_height()/2, f"{v:+.2f}", va="center", ha="left" if v>=0 else "right", fontsize=8, color=INK)
    finish(fig, "rate_sensitivity")

    fig, ax = plt.subplots(figsize=(7.0, 2.15))
    risks = ["ST LAP vintage / collections", "Gold collateral / custody", "ALM / funding maturity", "Cyber / data / systems", "Customer conduct / compliance"]
    levels = [3, 3, 3, 2, 2]
    cdict = {1:"#A9C9B5", 2:"#F0C46A", 3:"#D88962"}
    ax.barh(np.arange(len(risks)), levels, color=[cdict[i] for i in levels], height=.55)
    ax.set_yticks(np.arange(len(risks)), risks)
    ax.invert_yaxis()
    ax.set_xlim(0, 3.8)
    ax.set_xticks([1,2,3], ["Routine", "Ongoing", "High follow-up"])
    ax.set_title("External Analyst Follow-up Priority (Qualitative; Not a Fedfina Rating)", loc="left", color=NAVY)
    ax.grid(axis="x", color="#E5EAF0", linewidth=.75)
    ax.set_axisbelow(True)
    for side in ("top", "right", "left"):
        ax.spines[side].set_visible(False)
    ax.spines["bottom"].set_color("#AAB5BF")
    ax.tick_params(axis="x", labelsize=7.2, colors=MUTED, length=0)
    ax.tick_params(axis="y", labelsize=7.5, colors=INK, length=0)
    finish(fig, "risk_heatmap")

    return charts


# ---------- ReportLab PDF ----------

def register_pdf_fonts():
    font_dir = Path("/usr/share/fonts/truetype/dejavu")
    pdfmetrics.registerFont(TTFont("DVS", str(font_dir / "DejaVuSerif.ttf")))
    pdfmetrics.registerFont(TTFont("DVSB", str(font_dir / "DejaVuSerif-Bold.ttf")))
    pdfmetrics.registerFont(TTFont("DVSBI", str(font_dir / "DejaVuSerif-Bold.ttf")))
    pdfmetrics.registerFont(TTFont("DVSANS", str(font_dir / "DejaVuSans.ttf")))
    pdfmetrics.registerFont(TTFont("DVSANSB", str(font_dir / "DejaVuSans-Bold.ttf")))


def pdf_styles():
    return {
        "body": ParagraphStyle("body", fontName="DVS", fontSize=9.1, leading=12.4,
                               textColor=colors.HexColor(INK), spaceAfter=7, alignment=TA_JUSTIFY),
        "bullet": ParagraphStyle("bullet", fontName="DVS", fontSize=8.8, leading=11.7,
                                 textColor=colors.HexColor(INK), leftIndent=12, firstLineIndent=-8,
                                 bulletIndent=0, spaceAfter=3),
        "cell": ParagraphStyle("cell", fontName="DVS", fontSize=7.05, leading=8.9,
                               textColor=colors.HexColor(INK), spaceAfter=0),
        "headcell": ParagraphStyle("headcell", fontName="DVSANSB", fontSize=7.0, leading=8.6,
                                   textColor=colors.white, spaceAfter=0),
        "note": ParagraphStyle("note", fontName="DVSANS", fontSize=7.8, leading=10.2,
                               textColor=colors.HexColor("#5B4A26"), spaceAfter=0),
        "source": ParagraphStyle("source", fontName="DVSANS", fontSize=6.7, leading=8.4,
                                 textColor=colors.HexColor(MUTED), spaceAfter=0),
        "front": ParagraphStyle("front", fontName="DVS", fontSize=10.0, leading=16,
                                textColor=colors.HexColor("#111111"), alignment=TA_JUSTIFY, spaceAfter=14),
        "front_center": ParagraphStyle("front_center", fontName="DVS", fontSize=10.0, leading=15,
                                       textColor=colors.HexColor("#111111"), alignment=TA_CENTER, spaceAfter=10),
    }


def pdf_table(block, width, styles):
    rows = [block["headers"]] + block["rows"]
    ncol = len(block["headers"])
    if ncol == 2:
        widths = [width * .29, width * .71]
    elif ncol == 3:
        widths = [width * .25, width * .35, width * .40]
    elif ncol == 4:
        widths = [width * .23, width * .24, width * .23, width * .30]
    elif ncol == 5:
        widths = [width * .19, width * .21, width * .19, width * .20, width * .21]
    elif ncol == 7:
        widths = [width * .15, width * .13, width * .10, width * .13, width * .12, width * .17, width * .20]
    else:
        widths = [width/ncol] * ncol
    data = []
    for ridx, row in enumerate(rows):
        style = styles["headcell"] if ridx == 0 else styles["cell"]
        data.append([Paragraph(html.escape(str(value)).replace("\n", "<br/>"), style) for value in row])
    t = Table(data, colWidths=widths, repeatRows=1, hAlign="LEFT")
    t.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (-1,0), colors.HexColor(NAVY)),
        ("TEXTCOLOR", (0,0), (-1,0), colors.white),
        ("VALIGN", (0,0), (-1,-1), "TOP"),
        ("GRID", (0,0), (-1,-1), .35, colors.HexColor(GRID)),
        ("ROWBACKGROUNDS", (0,1), (-1,-1), [colors.white, colors.HexColor("#F5F8FA")]),
        ("LEFTPADDING", (0,0), (-1,-1), 4), ("RIGHTPADDING", (0,0), (-1,-1), 4),
        ("TOPPADDING", (0,0), (-1,-1), 4), ("BOTTOMPADDING", (0,0), (-1,-1), 4),
    ]))
    return t


def draw_border(c, w, h):
    c.setStrokeColor(colors.HexColor("#222222"))
    c.setLineWidth(1.4)
    c.rect(17, 17, w-34, h-34)
    c.setLineWidth(.6)
    c.rect(21, 21, w-42, h-42)


def draw_front_page(c, page_no, crest, letterhead, styles):
    w, h = A4
    draw_border(c, w, h)
    c.setFillColor(colors.HexColor("#111111"))
    if page_no == 1:
        y = 710
        for text, size, bold, gap in [
            ("SUMMER INTERNSHIP REPORT ON", 13, True, 30),
            ("RISK ANALYSIS AT FEDBANK FINANCIAL SERVICES LTD", 12, True, 26),
            ("FY 2025–26 COMPANY-DATA FOCUS", 9.2, True, 22),
            ("Submitted in partial fulfilment for the award of the degree of", 10, False, 22),
            ("MASTER OF MANAGEMENT STUDIES", 11, True, 18),
            ("FINANCE", 11, True, 40),
            ("UNIVERSITY OF MUMBAI", 11, True, 42),
            ("Submitted by", 10, True, 18),
            ("DATTAGURU PATIL", 11, True, 17),
            ("STUDENT OF SYMMS  |  ROLL NO. M16144", 10, True, 34),
            ("2026–2027", 10, True, 42),
            ("Under the Guidance of", 10, True, 18),
            ("DR. RAHUL N. WADEKAR", 10, True, 24),
        ]:
            c.setFont("DVSB" if bold else "DVS", size)
            c.drawCentredString(w/2, y, text)
            y -= gap
        c.drawImage(ImageReader(str(crest)), w/2-31, 181, width=62, height=62, preserveAspectRatio=True, mask="auto")
        c.setFont("DVSB", 9.1)
        c.drawCentredString(w/2, 155, "DES’S NAVINCHANDRA MEHTA INSTITUTE OF TECHNOLOGY")
        c.drawCentredString(w/2, 141, "AND DEVELOPMENT, MUMBAI – 400028")
    elif page_no == 2:
        c.drawImage(ImageReader(str(letterhead)), 49, 710, width=w-98, height=91, preserveAspectRatio=True, mask="auto")
        c.setFont("DVSB", 13); c.drawCentredString(w/2, 670, "CERTIFICATE")
        c.setFont("DVSANSB", 7.5); c.setFillColor(colors.HexColor("#9A5D12"))
        c.drawCentredString(w/2, 651, "DRAFT FOR INSTITUTIONAL VERIFICATION — SIGNATURES TO BE COMPLETED")
        text = ("This is to certify that the summer internship project report titled “Risk Analysis at Fedbank Financial Services Ltd” "
                "submitted by Dattaguru Patil, Roll No. M16144, student of S.Y. MMS (SYMMS), in partial fulfilment of the requirements "
                "for the MMS degree of the University of Mumbai, records a study of the internship work stated as Risk Analyst Intern "
                "at Fedbank Financial Services Ltd from 4 April 2026 to 3 July 2026. The report combines student-supplied internship "
                "particulars with analysis of the company’s public annual reports. The institution and project guide should review and "
                "authenticate this certificate wording before signature.")
        p = Paragraph(html.escape(text), styles["front"])
        pw, ph = p.wrap(w-140, 190); p.drawOn(c, 70, 620-ph)
        # Stamp area and signature placeholders, clearly unsigned.
        c.setStrokeColor(colors.HexColor("#AAB2B9")); c.setLineWidth(.7)
        c.rect(w/2-42, 337, 84, 75)
        c.setFont("DVSANS", 7); c.setFillColor(colors.HexColor(MUTED)); c.drawCentredString(w/2, 376, "COLLEGE STAMP")
        c.drawCentredString(w/2, 364, "(if required)")
        c.setFillColor(colors.HexColor("#111111")); c.setFont("DVS", 8.4)
        c.line(82, 326, 231, 326); c.line(365, 326, 514, 326)
        c.drawString(84, 311, "Dr. Rasika Mallya — Director")
        c.drawString(367, 311, "Dr. Rahul N. Wadekar — Project Guide")
        c.setFont("DVS", 8); c.drawCentredString(w/2, 257, "External Examiner: __________________________________")
    elif page_no == 3:
        c.setFont("DVSANSB", 9.5); c.setFillColor(colors.HexColor(NAVY))
        c.drawCentredString(w/2, 750, "FEDBANK FINANCIAL SERVICES LTD")
        c.setFillColor(colors.HexColor("#111111")); c.setFont("DVSB", 13)
        c.drawCentredString(w/2, 715, "INTERNSHIP PARTICULARS")
        c.setFont("DVSANSB", 7.4); c.setFillColor(colors.HexColor("#9A5D12"))
        c.drawCentredString(w/2, 696, "STUDENT-PREPARED SUMMARY — NOT AN EMPLOYER-ISSUED CERTIFICATE")
        c.setFillColor(colors.HexColor("#111111"))
        rows = [
            ("Student", "Dattaguru Patil"), ("Programme", "S.Y. MMS (SYMMS)"),
            ("Roll number", "M16144"), ("Organisation", "Fedbank Financial Services Ltd"),
            ("Role", "Risk Analyst Intern"), ("Internship period", "4 April 2026 to 3 July 2026"),
        ]
        yy = 642
        for label, value in rows:
            c.setFont("DVSB", 9.3); c.drawString(83, yy, label + ":")
            c.setFont("DVS", 9.3); c.drawString(205, yy, value)
            yy -= 29
        c.setFont("DVS", 9.1); c.drawString(83, yy-4, "Stated workstreams")
        tasks = [
            "Monthly preparation of performance graphs and EWS graphs (term to confirm).",
            "Quarterly preparation support for the Risk Management Committee (RMC).",
            "Monthly risk paper preparation and portfolio commentary.",
            "Risk-policy trigger versus actual-performance monitoring.",
            "Bounce analysis using approved presentment and return-code data.",
        ]
        yy -= 31
        for task in tasks:
            p = Paragraph("•  " + html.escape(task), ParagraphStyle("task", fontName="DVS", fontSize=8.8, leading=13))
            pw, ph = p.wrap(w-150, 60); p.drawOn(c, 91, yy-ph); yy -= ph+8
        box_y = 177
        c.setFillColor(colors.HexColor("#F5F7F8")); c.setStrokeColor(colors.HexColor(GRID))
        c.roundRect(70, box_y, w-140, 76, 5, stroke=1, fill=1)
        p = Paragraph(html.escape("This page records information supplied by the student. It does not certify employment, performance or completion on behalf of Fedbank Financial Services Ltd. Attach an original employer-issued certificate if required by the institute."), styles["note"])
        pw, ph = p.wrap(w-165, 60); p.drawOn(c, 82, box_y+58-ph)
    elif page_no == 4:
        c.setFont("DVSB", 13); c.drawCentredString(w/2, 650, "DECLARATION")
        text = ("I hereby declare that this project report, submitted in partial fulfilment of the requirements for the award of "
                "Master of Management Studies (MMS) of the University of Mumbai, is my academic work. It is based on the internship "
                "particulars stated by me and on the public annual reports identified in the bibliography. The analysis does not reproduce "
                "confidential borrower records, internal Fedfina policy thresholds or unpublished committee material. It has not been "
                "submitted to another University or Institute for the award of any degree, diploma or certificate. Any figures, calculations "
                "and interpretations derived from company disclosures are identified as such, and limitations are stated. I will verify all "
                "student, institutional and internship details before final submission.")
        p = Paragraph(html.escape(text), styles["front"])
        pw, ph = p.wrap(w-130, 250); p.drawOn(c, 65, 616-ph)
        c.setFont("DVS", 9.2)
        c.drawString(80, 255, "Name: Dattaguru Patil")
        c.drawString(80, 230, "Roll No.: M16144")
        c.drawString(365, 255, "Signature: ____________________")
        c.drawString(365, 230, "Date: ________________________")
    elif page_no == 5:
        c.setFont("DVSB", 13); c.drawCentredString(w/2, 690, "ACKNOWLEDGEMENT")
        text = ("This project has been a valuable opportunity to connect management studies with the practical discipline of risk analysis. "
                "I am grateful to my project guide, Dr. Rahul N. Wadekar, and to the faculty and staff of DES’s Navinchandra Mehta Institute "
                "of Technology and Development for their guidance and academic support. I thank Fedbank Financial Services Ltd and the "
                "colleagues who supported my Risk Analyst internship and introduced me to monthly portfolio reporting, quarterly RMC "
                "preparation, policy monitoring and bounce analysis. I also acknowledge the reporting teams whose annual reports made "
                "the company's public disclosures available for this study. I am thankful to my family and friends for their encouragement. "
                "Any errors or interpretations in this academic report remain my responsibility.")
        p = Paragraph(html.escape(text), styles["front_center"])
        pw, ph = p.wrap(w-90, 400); p.drawOn(c, 45, 662-ph)
    elif page_no == 6:
        c.setFont("DVSB", 13); c.drawCentredString(w/2, 710, "EXECUTIVE SUMMARY")
        paras = [
            "Summer internship training connects management concepts with workplace analysis. This project, “Risk Analysis at Fedbank Financial Services Ltd”, studies the principal risks of a retail-lending NBFC and the workstreams stated for my Risk Analyst internship from 4 April to 3 July 2026.",
            "The report focuses on FY2025–26, year ended 31 March 2026, and uses the FY2023–24 and FY2024–25 annual reports for historical context. Fedfina reported AUM of ₹20,153 crore, secured AUM of 98.9%, GNPA of 1.9%, NNPA of 1.3%, credit cost of 0.8%, PAT of ₹343.6 crore, ROA of 2.4% and ROE of 12.6%. Gold AUM reached ₹10,352 crore and the branch network stood at 757 across 17 states/UTs. These figures are company disclosures, not internal monthly internship results.",
            "The analysis reviews credit and collateral risk, asset quality, capital, funding, liquidity and maturity, interest-rate sensitivity, operations, cyber risk and customer conduct. It identifies FY26 growth and secured-mix expansion as constructive, while recommending close follow-up on ST LAP vintages, Gold custody/LTV, collections sustainability, longer-tenor liquidity, funding repricing, NNPA and reserve coverage.",
            "The stated tasks—monthly graphs and EWS graphs, quarterly RMC preparation, a monthly risk paper, policy triggers versus actuals and bounce analysis—are translated into reusable workflows and appendices. The task-list word “ever” is provisionally interpreted as EWS and requires supervisor confirmation. No unpublished policy limits, borrower records, bounce statistics or RMC material were supplied; templates are left blank rather than populated with invented values.",
        ]
        y = 678
        for text in paras:
            p = Paragraph(html.escape(text), styles["front_center"])
            pw, ph = p.wrap(w-92, 155); p.drawOn(c, 46, y-ph); y -= ph+9
        c.setFont("DVSANS", 7.6); c.setFillColor(colors.HexColor(MUTED))
        c.drawCentredString(w/2, 63, "Prepared 4 October 2026 · Public-source academic analysis · Subject to institutional review")
    c.setFillColor(colors.HexColor("#111111")); c.setFont("DVSANS", 7.4)
    c.drawString(34, 29, str(page_no))


def build_pdf(charts: dict[str, Path], crest: Path, letterhead: Path):
    register_pdf_fonts()
    styles = pdf_styles()
    w, h = A4
    c = canvas.Canvas(str(PDF_OUT), pagesize=A4, pageCompression=1)
    c.setTitle("Risk Analysis at Fedbank Financial Services Ltd")
    c.setAuthor("Dattaguru Patil")
    c.setSubject("Summer Internship Project Report · FY2025–26")
    c.setKeywords("Fedbank Financial Services, Fedfina, risk analysis, summer internship, FY2026")
    for page_no in range(1, 7):
        draw_front_page(c, page_no, crest, letterhead, styles)
        c.showPage()

    left, right = 48, 48
    content_w = w - left - right
    bottom = 50
    for page_no, spec in enumerate(PAGES, start=7):
        c.setFillColor(colors.HexColor(MUTED)); c.setFont("DVSANS", 7.0)
        c.drawString(left, h-27, "SUMMER INTERNSHIP REPORT  |  FY2025–26")
        c.drawRightString(w-right, h-27, "DATTAGURU PATIL  ·  M16144")
        c.setStrokeColor(colors.HexColor(NAVY)); c.setLineWidth(.8)
        c.line(left, h-34, w-right, h-34)
        c.setFillColor(colors.HexColor(ORANGE)); c.setFont("DVSANSB", 7.1)
        c.drawString(left, h-54, spec["section"].upper())
        title = Paragraph(html.escape(spec["title"]), ParagraphStyle("title", fontName="DVSB", fontSize=17.2, leading=20.5,
                                                                       textColor=colors.HexColor(NAVY), spaceAfter=0))
        tw, th = title.wrap(content_w, 48)
        title.drawOn(c, left, h-61-th)
        y = h-68-th-12
        overflows = []

        for text in spec["paragraphs"]:
            para = Paragraph(html.escape(text), styles["body"])
            pw, ph = para.wrap(content_w, max(20, y-bottom))
            if y-ph < bottom:
                overflows.append("paragraph")
            para.drawOn(c, left, y-ph); y -= ph + 3

        for bullet in spec["bullets"]:
            para = Paragraph(html.escape(bullet), styles["bullet"], bulletText="•")
            pw, ph = para.wrap(content_w, max(20, y-bottom))
            if y-ph < bottom:
                overflows.append("bullet")
            para.drawOn(c, left+4, y-ph); y -= ph + 2

        if spec["chart"]:
            img = RLImage(str(charts[spec["chart"]]))
            img._restrictSize(content_w, 150 if spec.get("table") else 177)
            iw, ih = img.wrap(content_w, max(20, y-bottom))
            if y-ih < bottom:
                overflows.append("chart")
            img.drawOn(c, left, y-ih); y -= ih + 7

        if spec["table"]:
            tab = pdf_table(spec["table"], content_w, styles)
            tw, th = tab.wrap(content_w, max(30, y-bottom))
            if y-th < bottom:
                overflows.append("table")
            tab.drawOn(c, left, y-th); y -= th + 7

        if spec["note"]:
            note = Paragraph(html.escape(spec["note"]), styles["note"])
            note_table = Table([[note]], colWidths=[content_w])
            note_table.setStyle(TableStyle([
                ("BACKGROUND", (0,0), (-1,-1), colors.HexColor("#FFF5DF")),
                ("BOX", (0,0), (-1,-1), .45, colors.HexColor("#E5C98E")),
                ("LEFTPADDING", (0,0), (-1,-1), 7), ("RIGHTPADDING", (0,0), (-1,-1), 7),
                ("TOPPADDING", (0,0), (-1,-1), 5), ("BOTTOMPADDING", (0,0), (-1,-1), 5),
            ]))
            nw, nh = note_table.wrap(content_w, max(20, y-bottom))
            if y-nh < bottom:
                overflows.append("note")
            note_table.drawOn(c, left, y-nh); y -= nh + 7

        if spec["sources"]:
            source_text = "Source: " + " ".join(spec["sources"])
            para = Paragraph(html.escape(source_text), styles["source"])
            sw, sh = para.wrap(content_w, max(14, y-bottom))
            if y-sh < bottom:
                overflows.append("source")
            para.drawOn(c, left, y-sh); y -= sh
        c.setFillColor(colors.HexColor(MUTED)); c.setFont("DVSANS", 7.2)
        c.drawCentredString(w/2, 25, str(page_no))
        if overflows:
            print(f"PDF layout warning p.{page_no} ({spec['title']}): {','.join(overflows)}; lowest y={y:.1f}")
        c.showPage()
    c.save()


# ---------- DOCX ----------
def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd"); tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=65, start=80, bottom=65, end=80):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar"); tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + m))
        if node is None:
            node = OxmlElement("w:" + m); tc_mar.append(node)
        node.set(qn("w:w"), str(v)); node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D5DEE5", size="5"):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders"); tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag); borders.append(el)
        el.set(qn("w:val"), "single"); el.set(qn("w:sz"), size)
        el.set(qn("w:space"), "0"); el.set(qn("w:color"), color)


def set_page_border(section):
    sect_pr = section._sectPr
    for existing in sect_pr.findall(qn("w:pgBorders")):
        sect_pr.remove(existing)
    borders = OxmlElement("w:pgBorders")
    borders.set(qn("w:offsetFrom"), "page")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement("w:" + edge)
        el.set(qn("w:val"), "double"); el.set(qn("w:sz"), "12")
        el.set(qn("w:space"), "11"); el.set(qn("w:color"), "222222")
        borders.append(el)
    sect_pr.append(borders)


def set_section_page_start(section, start):
    sect_pr = section._sectPr
    for n in sect_pr.findall(qn("w:pgNumType")):
        sect_pr.remove(n)
    node = OxmlElement("w:pgNumType"); node.set(qn("w:start"), str(start)); sect_pr.append(node)


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar"); fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar"); fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1); run._r.append(instr); run._r.append(fld_char2)
    run.font.name = "Times New Roman"; run.font.size = Pt(8); run.font.color.rgb = RGBColor(102, 115, 126)


def setup_docx_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"; normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor(32, 43, 53)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.05
    for name, size in (("Heading 1", 16), ("Heading 2", 12)):
        st = styles[name]
        st.font.name = "Times New Roman"; st.font.size = Pt(size); st.font.bold = True
        st.font.color.rgb = RGBColor(23, 70, 107)
        st.paragraph_format.space_before = Pt(4); st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.keep_with_next = True


def add_front_paragraph(doc, text, size=10, bold=False, before=0, after=8, align=WD_ALIGN_PARAGRAPH.CENTER):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(before); p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.05
    r = p.add_run(text)
    r.font.name = "Times New Roman"; r.font.size = Pt(size); r.bold = bold
    return p


def add_front_matter(doc, crest, letterhead):
    # Page 1: cover page patterned on the reference.
    add_front_paragraph(doc, "SUMMER INTERNSHIP REPORT ON", 12, True, before=38, after=23)
    add_front_paragraph(doc, "RISK ANALYSIS AT FEDBANK FINANCIAL SERVICES LTD", 12, True, after=7)
    add_front_paragraph(doc, "FY 2025–26 COMPANY-DATA FOCUS", 8.8, True, after=18)
    add_front_paragraph(doc, "Submitted in partial fulfilment for the award of the degree of", 10, False, after=15)
    add_front_paragraph(doc, "MASTER OF MANAGEMENT STUDIES", 11, True, after=3)
    add_front_paragraph(doc, "FINANCE", 11, True, after=23)
    add_front_paragraph(doc, "UNIVERSITY OF MUMBAI", 11, True, after=25)
    add_front_paragraph(doc, "Submitted by", 9.5, True, after=4)
    add_front_paragraph(doc, "DATTAGURU PATIL", 10.5, True, after=3)
    add_front_paragraph(doc, "STUDENT OF SYMMS  |  ROLL NO. M16144", 9.5, True, after=20)
    add_front_paragraph(doc, "2026–2027", 10, True, after=22)
    add_front_paragraph(doc, "Under the Guidance of", 9.5, True, after=4)
    add_front_paragraph(doc, "DR. RAHUL N. WADEKAR", 9.5, True, after=10)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(5)
    p.add_run().add_picture(str(crest), width=Inches(.68), height=Inches(.68))
    add_front_paragraph(doc, "DES’S NAVINCHANDRA MEHTA INSTITUTE OF TECHNOLOGY", 8.7, True, after=1)
    add_front_paragraph(doc, "AND DEVELOPMENT, MUMBAI – 400028", 8.7, True, after=0)
    doc.add_page_break()

    # Page 2: certificate layout, with unsigned/verification status explicit.
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(6)
    p.add_run().add_picture(str(letterhead), width=Inches(6.5))
    add_front_paragraph(doc, "CERTIFICATE", 12, True, before=0, after=5)
    add_front_paragraph(doc, "DRAFT FOR INSTITUTIONAL VERIFICATION — SIGNATURES TO BE COMPLETED", 7.5, True, after=13)
    text = ("This is to certify that the summer internship project report titled “Risk Analysis at Fedbank Financial Services Ltd” submitted by Dattaguru Patil, Roll No. M16144, student of S.Y. MMS (SYMMS), in partial fulfilment of the requirements for the MMS degree of the University of Mumbai, records a study of the internship work stated as Risk Analyst Intern at Fedbank Financial Services Ltd from 4 April 2026 to 3 July 2026. The report combines student-supplied internship particulars with analysis of the company’s public annual reports. The institution and project guide should review and authenticate this certificate wording before signature.")
    p = doc.add_paragraph(text); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Inches(.48); p.paragraph_format.right_indent = Inches(.48)
    p.paragraph_format.space_after = Pt(16); p.paragraph_format.line_spacing = 1.15
    for r in p.runs: r.font.name="Times New Roman"; r.font.size=Pt(9)
    t = doc.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    t.columns[0].width = Inches(1.2); t.rows[0].height = Inches(.75)
    cell=t.cell(0,0); cell.width=Inches(1.2); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_shading(cell,"FFFFFF"); set_cell_margins(cell,30,30,30,30); set_table_borders(t,"9DA8B1","5")
    p=cell.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    rr=p.add_run("COLLEGE STAMP\n(if required)"); rr.font.name="Times New Roman"; rr.font.size=Pt(7)
    doc.add_paragraph()
    sig = doc.add_table(rows=1, cols=2); sig.alignment=WD_TABLE_ALIGNMENT.CENTER; sig.autofit=False
    sig.columns[0].width=Inches(3.0); sig.columns[1].width=Inches(3.3)
    for cell, text in zip(sig.rows[0].cells,["____________________________\nDr. Rasika Mallya — Director", "____________________________\nDr. Rahul N. Wadekar — Project Guide"]):
        cell.text=text; set_cell_margins(cell,50,80,50,80)
        for p in cell.paragraphs:
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs: r.font.name="Times New Roman"; r.font.size=Pt(8)
    add_front_paragraph(doc, "External Examiner: __________________________________", 8, False, before=13, after=0)
    doc.add_page_break()

    # Page 3: neutral student-prepared particulars, not an employer certificate.
    add_front_paragraph(doc, "FEDBANK FINANCIAL SERVICES LTD", 9.5, True, before=26, after=19)
    add_front_paragraph(doc, "INTERNSHIP PARTICULARS", 12, True, after=5)
    add_front_paragraph(doc, "STUDENT-PREPARED SUMMARY — NOT AN EMPLOYER-ISSUED CERTIFICATE", 7.5, True, after=20)
    rows = [
        ("Student", "Dattaguru Patil"), ("Programme", "S.Y. MMS (SYMMS)"),
        ("Roll number", "M16144"), ("Organisation", "Fedbank Financial Services Ltd"),
        ("Role", "Risk Analyst Intern"), ("Internship period", "4 April 2026 to 3 July 2026"),
    ]
    for label, value in rows:
        p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(.95); p.paragraph_format.space_after=Pt(7)
        r=p.add_run(label+":  "); r.bold=True; r.font.name="Times New Roman"; r.font.size=Pt(9.5)
        r=p.add_run(value); r.font.name="Times New Roman"; r.font.size=Pt(9.5)
    p=doc.add_paragraph("Stated workstreams"); p.paragraph_format.left_indent=Inches(.95); p.paragraph_format.space_before=Pt(6)
    p.runs[0].bold=True; p.runs[0].font.size=Pt(9.5)
    tasks=["Monthly preparation of performance graphs and EWS graphs (term to confirm).", "Quarterly preparation support for the Risk Management Committee (RMC).", "Monthly risk paper preparation and portfolio commentary.", "Risk-policy trigger versus actual-performance monitoring.", "Bounce analysis using approved presentment and return-code data."]
    for task in tasks:
        p=doc.add_paragraph(style="List Bullet"); p.paragraph_format.left_indent=Inches(1.1); p.paragraph_format.first_line_indent=Inches(-.18); p.paragraph_format.space_after=Pt(4)
        p.add_run(task).font.size=Pt(8.8)
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(18); p.paragraph_format.left_indent=Inches(.6); p.paragraph_format.right_indent=Inches(.6)
    pPr=p._p.get_or_add_pPr(); shd=OxmlElement("w:shd"); shd.set(qn("w:fill"),"F5F7F8"); pPr.append(shd)
    r=p.add_run("This page records information supplied by the student. It does not certify employment, performance or completion on behalf of Fedbank Financial Services Ltd. Attach an original employer-issued certificate if required by the institute.")
    r.font.size=Pt(8); r.font.color.rgb=RGBColor(91,74,38)
    doc.add_page_break()

    # Page 4: declaration.
    add_front_paragraph(doc, "DECLARATION", 12, True, before=88, after=26)
    declaration=("I hereby declare that this project report, submitted in partial fulfilment of the requirements for the award of Master of Management Studies (MMS) of the University of Mumbai, is my academic work. It is based on the internship particulars stated by me and on the public annual reports identified in the bibliography. The analysis does not reproduce confidential borrower records, internal Fedfina policy thresholds or unpublished committee material. It has not been submitted to another University or Institute for the award of any degree, diploma or certificate. Any figures, calculations and interpretations derived from company disclosures are identified as such, and limitations are stated. I will verify all student, institutional and internship details before final submission.")
    p=doc.add_paragraph(declaration); p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent=Inches(.55); p.paragraph_format.right_indent=Inches(.55); p.paragraph_format.line_spacing=1.2; p.paragraph_format.space_after=Pt(15)
    for r in p.runs: r.font.size=Pt(9.5)
    for _ in range(5): doc.add_paragraph()
    t=doc.add_table(rows=2,cols=2); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    t.columns[0].width=Inches(3.1); t.columns[1].width=Inches(3.1)
    t.cell(0,0).text="Name: Dattaguru Patil"; t.cell(0,1).text="Signature: ____________________"
    t.cell(1,0).text="Roll No.: M16144"; t.cell(1,1).text="Date: ________________________"
    for row in t.rows:
        for cell in row.cells:
            set_cell_margins(cell,40,80,40,80)
            for p in cell.paragraphs:
                for r in p.runs: r.font.name="Times New Roman"; r.font.size=Pt(8.8)
    doc.add_page_break()

    # Page 5: acknowledgement.
    add_front_paragraph(doc, "ACKNOWLEDGEMENT", 12, True, before=54, after=29)
    ack=("This project has been a valuable opportunity to connect management studies with the practical discipline of risk analysis. I am grateful to my project guide, Dr. Rahul N. Wadekar, and to the faculty and staff of DES’s Navinchandra Mehta Institute of Technology and Development for their guidance and academic support. I thank Fedbank Financial Services Ltd and the colleagues who supported my Risk Analyst internship and introduced me to monthly portfolio reporting, quarterly RMC preparation, policy monitoring and bounce analysis. I also acknowledge the reporting teams whose annual reports made the company’s public disclosures available for this study. I am thankful to my family and friends for their encouragement. Any errors or interpretations in this academic report remain my responsibility.")
    p=doc.add_paragraph(ack); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.left_indent=Inches(.35); p.paragraph_format.right_indent=Inches(.35); p.paragraph_format.line_spacing=1.2
    for r in p.runs: r.font.size=Pt(9.7)
    doc.add_page_break()

    # Page 6: executive summary.
    add_front_paragraph(doc, "EXECUTIVE SUMMARY", 12, True, before=30, after=21)
    summary=[
        "Summer internship training connects management concepts with workplace analysis. This project, “Risk Analysis at Fedbank Financial Services Ltd”, studies the principal risks of a retail-lending NBFC and the workstreams stated for my Risk Analyst internship from 4 April to 3 July 2026.",
        "The report focuses on FY2025–26, year ended 31 March 2026, and uses the FY2023–24 and FY2024–25 annual reports for historical context. Fedfina reported AUM of ₹20,153 crore, secured AUM of 98.9%, GNPA of 1.9%, NNPA of 1.3%, credit cost of 0.8%, PAT of ₹343.6 crore, ROA of 2.4% and ROE of 12.6%. Gold AUM reached ₹10,352 crore and the branch network stood at 757 across 17 states/UTs. These figures are company disclosures, not internal monthly internship results.",
        "The analysis reviews credit and collateral risk, asset quality, capital, funding, liquidity and maturity, interest-rate sensitivity, operations, cyber risk and customer conduct. It identifies FY26 growth and secured-mix expansion as constructive, while recommending close follow-up on ST LAP vintages, Gold custody/LTV, collections sustainability, longer-tenor liquidity, funding repricing, NNPA and reserve coverage.",
        "The stated tasks—monthly graphs and EWS graphs, quarterly RMC preparation, a monthly risk paper, policy triggers versus actuals and bounce analysis—are translated into reusable workflows and appendices. The task-list word “ever” is provisionally interpreted as EWS and requires supervisor confirmation. No unpublished policy limits, borrower records, bounce statistics or RMC material were supplied; templates are left blank rather than populated with invented values.",
    ]
    for text in summary:
        p=doc.add_paragraph(text); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.left_indent=Inches(.28); p.paragraph_format.right_indent=Inches(.28); p.paragraph_format.line_spacing=1.05; p.paragraph_format.space_after=Pt(9)
        for r in p.runs: r.font.size=Pt(8.7)
    p=doc.add_paragraph("Prepared 4 October 2026 · Public-source academic analysis · Subject to institutional review")
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(4)
    for r in p.runs: r.font.size=Pt(7.2); r.font.color.rgb=RGBColor(102,115,126)


def add_docx_table(doc, block):
    rows = [block["headers"]] + block["rows"]
    ncol = len(block["headers"])
    table = doc.add_table(rows=len(rows), cols=ncol)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER; table.autofit = False
    table.style = "Table Grid"
    avail = 6.85
    if ncol == 2: weights=[.29,.71]
    elif ncol == 3: weights=[.25,.35,.40]
    elif ncol == 4: weights=[.23,.24,.23,.30]
    elif ncol == 5: weights=[.19,.21,.19,.20,.21]
    elif ncol == 7: weights=[.15,.13,.10,.13,.12,.17,.20]
    else: weights=[1/ncol]*ncol
    for col,wgt in enumerate(weights):
        table.columns[col].width=Inches(avail*wgt)
    for rix,row in enumerate(rows):
        for cix,value in enumerate(row):
            cell=table.cell(rix,cix); cell.width=Inches(avail*weights[cix]); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
            set_cell_margins(cell,45,55,45,55)
            if rix==0: set_cell_shading(cell,"17466B")
            elif rix % 2 == 0: set_cell_shading(cell,"F5F8FA")
            p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.0
            run=p.add_run(str(value)); run.font.name="Times New Roman"; run.font.size=Pt(7.3 if ncol>=5 else 7.8)
            run.bold=(rix==0); run.font.color.rgb=RGBColor(255,255,255) if rix==0 else RGBColor(32,43,53)
    set_table_borders(table)
    return table


def add_docx_note(doc, text, source=False):
    table=doc.add_table(rows=1,cols=1); table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.autofit=False
    table.columns[0].width=Inches(6.85)
    cell=table.cell(0,0); cell.width=Inches(6.85); set_cell_margins(cell,45,75,45,75)
    set_cell_shading(cell,"FFF5DF" if not source else "F4F6F8")
    set_table_borders(table,"E5C98E" if not source else "E2E7EB","4")
    p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1
    r=p.add_run(text); r.font.name="Times New Roman"; r.font.size=Pt(7.6 if not source else 6.8)
    r.font.color.rgb=RGBColor(91,74,38) if not source else RGBColor(102,115,126)


def configure_section(section, body=False):
    section.page_width=Inches(8.27); section.page_height=Inches(11.69)
    section.left_margin=Inches(.70); section.right_margin=Inches(.70)
    section.top_margin=Inches(.63 if body else .55); section.bottom_margin=Inches(.58)
    section.header_distance=Inches(.28); section.footer_distance=Inches(.27)


def set_body_header_footer(section):
    header=section.header
    p=header.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after=Pt(1)
    r=p.add_run("SUMMER INTERNSHIP REPORT  |  RISK ANALYSIS AT FEDFINA  |  FY2025–26")
    r.font.name="Times New Roman"; r.font.size=Pt(7.2); r.font.color.rgb=RGBColor(102,115,126)
    pPr=p._p.get_or_add_pPr(); pBdr=OxmlElement("w:pBdr"); bottom=OxmlElement("w:bottom")
    bottom.set(qn("w:val"),"single"); bottom.set(qn("w:sz"),"5"); bottom.set(qn("w:space"),"2"); bottom.set(qn("w:color"),NAVY.replace("#","")); pBdr.append(bottom); pPr.append(pBdr)
    footer=section.footer
    p=footer.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    add_page_field(p)


def build_docx(charts: dict[str, Path], crest: Path, letterhead: Path):
    doc=Document()
    setup_docx_styles(doc)
    configure_section(doc.sections[0], body=False)
    set_page_border(doc.sections[0]); set_section_page_start(doc.sections[0],1)
    doc.sections[0].header.paragraphs[0].text=""
    footer=doc.sections[0].footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.LEFT
    add_page_field(footer)
    add_front_matter(doc, crest, letterhead)

    body_section=doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(body_section, body=True)
    body_section.header.is_linked_to_previous = False
    body_section.footer.is_linked_to_previous = False
    # Remove page border from body and continue page numbering from 7.
    for node in body_section._sectPr.findall(qn("w:pgBorders")):
        body_section._sectPr.remove(node)
    set_section_page_start(body_section,7)
    set_body_header_footer(body_section)

    for page_index,spec in enumerate(PAGES, start=7):
        p=doc.add_paragraph(spec["section"].upper())
        p.paragraph_format.space_after=Pt(1); p.paragraph_format.keep_with_next=True
        r=p.runs[0]; r.font.name="Times New Roman"; r.font.size=Pt(7.5); r.bold=True; r.font.color.rgb=RGBColor(232,162,58)
        p=doc.add_paragraph(spec["title"],style="Heading 1")
        p.paragraph_format.space_before=Pt(0); p.paragraph_format.space_after=Pt(5); p.paragraph_format.keep_with_next=True
        for text in spec["paragraphs"]:
            p=doc.add_paragraph(text)
            p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.space_after=Pt(4); p.paragraph_format.line_spacing=1.02
            for r in p.runs: r.font.name="Times New Roman"; r.font.size=Pt(8.7)
        for bullet in spec["bullets"]:
            p=doc.add_paragraph(style="List Bullet")
            p.paragraph_format.left_indent=Inches(.22); p.paragraph_format.first_line_indent=Inches(-.12)
            p.paragraph_format.space_after=Pt(2); p.paragraph_format.line_spacing=1.0
            r=p.add_run(bullet); r.font.name="Times New Roman"; r.font.size=Pt(8.2)
        if spec["chart"]:
            p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before=Pt(2); p.paragraph_format.space_after=Pt(3)
            p.add_run().add_picture(str(charts[spec["chart"]]), width=Inches(6.42))
        if spec["table"]:
            add_docx_table(doc,spec["table"])
            p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1
        if spec["note"]:
            add_docx_note(doc,spec["note"])
        if spec["sources"]:
            add_docx_note(doc,"Source: " + " ".join(spec["sources"]),source=True)
        if page_index < 100:
            doc.add_page_break()

    props=doc.core_properties
    props.title="Risk Analysis at Fedbank Financial Services Ltd"
    props.subject="Summer Internship Project Report · FY2025–26"
    props.author="Dattaguru Patil"
    props.keywords="Fedfina, risk analysis, summer internship, FY2026"
    props.comments="Prepared from supplied annual reports; student-prepared internship particulars."
    doc.save(DOCX_OUT)


# ---------- Companion workbook ----------

def style_sheet(ws, title_row=1, header_row=3):
    ws.sheet_view.showGridLines=False
    ws.freeze_panes=f"A{header_row+1}"
    ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0
    ws.sheet_view.zoomScale=90
    for cell in ws[title_row]:
        if cell.value is not None:
            cell.font=Font(name="Aptos Display", size=15, bold=True, color="FFFFFF")
            cell.fill=PatternFill("solid",fgColor="17466B")
            cell.alignment=Alignment(vertical="center")
    ws.row_dimensions[title_row].height=27
    for cell in ws[header_row]:
        if cell.value is not None:
            cell.font=Font(name="Aptos", size=10, bold=True, color="FFFFFF")
            cell.fill=PatternFill("solid",fgColor="2F759D")
            cell.alignment=Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[header_row].height=34
    thin=Side(style="thin",color="D5DEE5")
    for row in ws.iter_rows(min_row=header_row+1):
        for cell in row:
            cell.border=Border(bottom=thin)
            cell.alignment=Alignment(vertical="top",wrap_text=True)
            if cell.value is not None:
                cell.font=Font(name="Aptos", size=9, color="25323A")


def build_workbook():
    wb=Workbook()
    ws=wb.active; ws.title="Start Here"
    ws.sheet_view.showGridLines=False
    ws.merge_cells("A1:F1"); ws["A1"]="Fedfina Risk Analyst Internship Workbook"
    ws["A1"].font=Font(name="Aptos Display",size=16,bold=True,color="FFFFFF"); ws["A1"].fill=PatternFill("solid",fgColor="17466B")
    ws.row_dimensions[1].height=30
    intro=[
        ("Purpose","Blank academic templates for monthly graphs/EWS, policy triggers, bounce analysis, risk papers and quarterly RMC preparation."),
        ("Company-data year","FY2025–26 annual-report trend data is included on the FY26 Trend tab; FY26 is the primary analytical period."),
        ("No invented MIS","Monthly, bounce and internal policy values are intentionally blank. No borrower-level data or internal thresholds were supplied."),
        ("Confidentiality","Populate only in an authorised environment with approved, minimised data. Do not commit confidential workbooks to Git."),
        ("Metric definitions","Confirm source, cut-off, denominator, retry rules and policy version before use. The approved Fedfina definitions take precedence."),
        ("Workbook limitation","Formulas are templates and must be checked against company systems and reporting rules before operational use."),
        ("Student","Dattaguru Patil · Roll No. M16144 · Risk Analyst Intern · 4 Apr–3 Jul 2026"),
    ]
    ws.append(["Topic","Guidance"])
    for row in intro: ws.append(list(row))
    ws.column_dimensions["A"].width=24; ws.column_dimensions["B"].width=105
    ws.freeze_panes="A3"
    for cell in ws[2]: cell.font=Font(bold=True,color="FFFFFF"); cell.fill=PatternFill("solid",fgColor="2F759D")
    for r in range(3,10):
        ws.cell(r,1).font=Font(bold=True,color="17466B")
        ws.cell(r,2).alignment=Alignment(wrap_text=True,vertical="top")
        ws.row_dimensions[r].height=33
    ws.sheet_view.zoomScale=90

    trend=wb.create_sheet("FY26 Trend")
    trend.append(["Company-Reported Five-Year Trend | ₹ crore unless noted"])
    trend.append(["Source: Fedfina Annual Report 2025–26, supplied PDF pp.14–15. Latest series used consistently; historical definitions may differ from prior reports."])
    trend.append(["Metric","Unit","FY22","FY23","FY24","FY25","FY26","FY26 YoY"])
    rows=[
        ["AUM","₹ Cr",*AUM_CR], ["PAT","₹ Cr",*PAT_CR], ["Yield","%",*YIELD], ["Spread","%",*SPREAD],
        ["Cost of borrowings","%",*COF], ["Cost-to-income","%",*COST_INCOME], ["ROA","%",*ROA], ["ROE","%",*ROE],
        ["GNPA","%",*GNPA], ["NNPA","%",*NNPA], ["Provision coverage","%",*PCR], ["CRAR","%",*CRAR],
    ]
    for row in rows:
        trend.append(row+[f'=IFERROR((G{trend.max_row+1}/F{trend.max_row+1})-1,"")'])
    style_sheet(trend,1,3)
    trend.merge_cells("A1:H1"); trend.merge_cells("A2:H2")
    trend["A2"].font=Font(name="Aptos",size=9,italic=True,color="66737E"); trend["A2"].alignment=Alignment(wrap_text=True)
    trend.row_dimensions[2].height=30
    trend.column_dimensions["A"].width=25; trend.column_dimensions["B"].width=13
    for col in range(3,9): trend.column_dimensions[get_column_letter(col)].width=13
    for r in range(4,trend.max_row+1):
        if trend.cell(r,2).value=="%":
            for c in range(3,8): trend.cell(r,c).number_format="0.0"
            trend.cell(r,8).number_format="0.0%"
        else:
            for c in range(3,8): trend.cell(r,c).number_format="#,##0.0"
            trend.cell(r,8).number_format="0.0%"
    # Charts linked to the public annual trend data.
    ch=LineChart(); ch.title="AUM trend (₹ Cr)"; ch.y_axis.title="₹ crore"; ch.x_axis.title="Financial year"; ch.height=7; ch.width=14
    ch.add_data(Reference(trend,min_col=3,max_col=7,min_row=4,max_row=4),from_rows=True,titles_from_data=False)
    ch.set_categories(Reference(trend,min_col=3,max_col=7,min_row=3))
    ch.series[0].graphicalProperties.line.solidFill="17466B"; ch.series[0].graphicalProperties.line.width=26000
    trend.add_chart(ch,"A18")
    ch2=LineChart(); ch2.title="Asset quality trend (%)"; ch2.y_axis.title="Percent"; ch2.height=7; ch2.width=14
    ch2.add_data(Reference(trend,min_col=3,max_col=7,min_row=12,max_row=13),from_rows=True,titles_from_data=False)
    ch2.set_categories(Reference(trend,min_col=3,max_col=7,min_row=3))
    trend.add_chart(ch2,"H18")

    monthly=wb.create_sheet("Monthly Data")
    headers=["Month","Product / segment","AUM (₹ Cr)","Disbursals (₹ Cr)","Collections due (₹ Cr)","Collections received (₹ Cr)","Valid presentments (#)","Bounced presentments (#)","Amount presented (₹)","Amount returned (₹)","GNPA (%)","NNPA (%)","Stage 2 (%)","Cost of funds (%)","Credit cost (%)","Collection efficiency (%)","Count bounce rate (%)","Amount bounce rate (%)","Source / owner","Comments"]
    monthly.append(["Monthly Risk Dashboard Input | Populate only from authorised MIS"])
    monthly.append(["Blank by design. Confirm definitions, period cut-off, retry handling, source owner and policy version before use."])
    monthly.append(headers)
    for r in range(4,40):
        monthly.cell(r,16,f'=IF(OR(E{r}="",F{r}=""),"",IFERROR(F{r}/E{r},""))')
        monthly.cell(r,17,f'=IF(OR(G{r}="",H{r}=""),"",IFERROR(H{r}/G{r},""))')
        monthly.cell(r,18,f'=IF(OR(I{r}="",J{r}=""),"",IFERROR(J{r}/I{r},""))')
        for c in range(11,19): monthly.cell(r,c).number_format="0.0%"
    style_sheet(monthly,1,3); monthly.merge_cells("A1:T1"); monthly.merge_cells("A2:T2")
    monthly["A2"].font=Font(name="Aptos",size=9,italic=True,color="66737E"); monthly["A2"].alignment=Alignment(wrap_text=True)
    widths=[15,19,15,17,19,21,20,21,20,19,13,13,13,17,15,22,19,21,20,34]
    for i,width in enumerate(widths,1): monthly.column_dimensions[get_column_letter(i)].width=width
    product_dv=DataValidation(type="list",formula1='"Gold,ST LAP,MT LAP,Home Loan,Other"',allow_blank=True)
    monthly.add_data_validation(product_dv); product_dv.add("B4:B39")
    # monthly trend charts: blank until actual values are entered.
    c1=LineChart(); c1.title="Monthly AUM by segment (₹ Cr)"; c1.y_axis.title="₹ Cr"; c1.x_axis.title="Month"; c1.height=7; c1.width=14
    c1.add_data(Reference(monthly,min_col=3,max_col=3,min_row=3,max_row=39),titles_from_data=True); c1.set_categories(Reference(monthly,min_col=1,min_row=4,max_row=39)); monthly.add_chart(c1,"A42")
    c2=LineChart(); c2.title="Bounce rates (monthly)"; c2.y_axis.title="Rate"; c2.x_axis.title="Month"; c2.height=7; c2.width=14
    c2.add_data(Reference(monthly,min_col=17,max_col=18,min_row=3,max_row=39),titles_from_data=True); c2.set_categories(Reference(monthly,min_col=1,min_row=4,max_row=39)); monthly.add_chart(c2,"J42")

    triggers=wb.create_sheet("Policy Triggers")
    triggers.append(["Policy Trigger vs Actual | Internal limits intentionally blank"])
    triggers.append(["Use current approved policy values and direction. MAX: breach if actual exceeds limit. MIN: breach if actual falls below limit. Preserve every breach and closure event."])
    triggers.append(["Metric / policy clause","Direction (MAX/MIN)","Approved limit","Actual","As of date","Status","Headroom","Owner","Root cause / action","Due date","Closure evidence"])
    for r in range(4,24):
        triggers.cell(r,6,f'=IF(OR(A{r}="",B{r}="",C{r}="",D{r}=""),"DATA NEEDED",IF(B{r}="MAX",IF(D{r}>C{r},"BREACH","WITHIN"),IF(B{r}="MIN",IF(D{r}<C{r},"BREACH","WITHIN"),"CHECK DIRECTION")))')
        triggers.cell(r,7,f'=IF(OR(C{r}="",D{r}=""),"",IF(B{r}="MAX",C{r}-D{r},D{r}-C{r}))')
    style_sheet(triggers,1,3); triggers.merge_cells("A1:K1"); triggers.merge_cells("A2:K2")
    triggers["A2"].font=Font(name="Aptos",size=9,italic=True,color="66737E"); triggers["A2"].alignment=Alignment(wrap_text=True); triggers.row_dimensions[2].height=32
    for i,width in enumerate([28,19,17,15,16,19,15,20,38,15,34],1): triggers.column_dimensions[get_column_letter(i)].width=width
    dv=DataValidation(type="list",formula1='"MAX,MIN"',allow_blank=True); triggers.add_data_validation(dv); dv.add("B4:B23")
    triggers.conditional_formatting.add("F4:F23",FormulaRule(formula=['F4="BREACH"'],fill=PatternFill("solid",fgColor="F4CCCC"),font=Font(color="9C0006",bold=True)))
    triggers.conditional_formatting.add("F4:F23",FormulaRule(formula=['F4="WITHIN"'],fill=PatternFill("solid",fgColor="D9EAD3"),font=Font(color="274E13")))

    bounce=wb.create_sheet("Bounce Analysis")
    bounce.append(["Bounce Analysis | Blank template — not company results"])
    bounce.append(["Define valid presentments, retry treatment, return-code taxonomy and cure window before calculating any rate."])
    bounce.append(["Month","Product / segment","Presentments (#)","Bounced (#)","Amount presented (₹)","Amount returned (₹)","Count bounce rate","Amount bounce rate","Cured within window (#)","Cure rate","Repeat-bounce accounts (#)","Repeat rate","Primary return reason","Source / owner","Notes"])
    for r in range(4,40):
        bounce.cell(r,7,f'=IF(OR(C{r}="",D{r}=""),"",IFERROR(D{r}/C{r},""))')
        bounce.cell(r,8,f'=IF(OR(E{r}="",F{r}=""),"",IFERROR(F{r}/E{r},""))')
        bounce.cell(r,10,f'=IF(OR(D{r}="",I{r}=""),"",IFERROR(I{r}/D{r},""))')
        bounce.cell(r,12,f'=IF(OR(C{r}="",K{r}=""),"",IFERROR(K{r}/C{r},""))')
        for c in (7,8,10,12): bounce.cell(r,c).number_format="0.0%"
    style_sheet(bounce,1,3); bounce.merge_cells("A1:O1"); bounce.merge_cells("A2:O2")
    bounce["A2"].font=Font(name="Aptos",size=9,italic=True,color="66737E"); bounce["A2"].alignment=Alignment(wrap_text=True); bounce.row_dimensions[2].height=31
    for i,width in enumerate([15,19,18,16,20,20,20,21,23,15,24,17,30,21,35],1): bounce.column_dimensions[get_column_letter(i)].width=width

    paper=wb.create_sheet("Monthly Risk Paper")
    paper.append(["Monthly Risk Paper | Working outline"])
    paper.append(["Populate using authorised sources. Separate fact, interpretation, action and decision request."])
    paper.append(["Section","Prompt","Reporting period","Source / owner","Reviewer comments","Status / action owner"])
    prompts=[
        ("Executive summary","Material movements, breaches, decisions required"),
        ("Portfolio / vintage","AUM, disbursals, product/vintage quality, concentration"),
        ("Credit / EWS","DPD, Stage movement, EWS alerts, cure and ECL"),
        ("Bounce / collections","Presentments, return codes, cure, recovery, conduct"),
        ("Funding / ALM","Maturity, liquidity, repricing, FX/hedge, concentration"),
        ("Operational / cyber","Incidents, audit findings, BCP, vendor and data quality"),
        ("Customer / compliance","Complaints, regulatory change, service and conduct"),
        ("Actions","Owner, due date, evidence, overdue items and escalation"),
    ]
    for sec,prompt in prompts: paper.append([sec,prompt,"","","",""])
    style_sheet(paper,1,3); paper.merge_cells("A1:F1"); paper.merge_cells("A2:F2")
    paper["A2"].font=Font(name="Aptos",size=9,italic=True,color="66737E"); paper["A2"].alignment=Alignment(wrap_text=True); paper.row_dimensions[2].height=30
    for i,width in enumerate([25,56,20,27,34,29],1): paper.column_dimensions[get_column_letter(i)].width=width
    for r in range(4,12): paper.row_dimensions[r].height=42

    rmc=wb.create_sheet("RMC Pack")
    rmc.append(["Quarterly RMC Pack | Preparation and Action Log"])
    rmc.append(["Template only. Confirm the current Risk Management Committee terms of reference and approved internal reporting requirements."])
    rmc.append(["Agenda item","Quarter / period","Key movement / issue","Evidence / source","Decision requested","Action owner","Due date","Status / closure evidence"])
    agenda=["Risk appetite / policy","Credit portfolio / product EWS","NPA / collections / bounce","Liquidity / ALM / funding","Capital / interest-rate / FX","Operational / fraud","Cyber / third party / BCP","Compliance / conduct / customer","Prior committee actions"]
    for item in agenda: rmc.append([item,"","","","","","",""])
    style_sheet(rmc,1,3); rmc.merge_cells("A1:H1"); rmc.merge_cells("A2:H2")
    rmc["A2"].font=Font(name="Aptos",size=9,italic=True,color="66737E"); rmc["A2"].alignment=Alignment(wrap_text=True); rmc.row_dimensions[2].height=32
    for i,width in enumerate([31,19,43,30,34,22,17,38],1): rmc.column_dimensions[get_column_letter(i)].width=width
    for r in range(4,13): rmc.row_dimensions[r].height=38

    # Global workbook tab color and properties.
    for sheet in wb.worksheets:
        sheet.sheet_properties.tabColor="17466B"
    wb.calculation.fullCalcOnLoad=True
    wb.calculation.forceFullCalc=True
    wb.calculation.calcMode="auto"
    wb.properties.title="Fedfina Risk Analyst Internship Workbook"
    wb.properties.creator="Dattaguru Patil"
    wb.properties.subject="FY2025-26 public trend series and blank risk-monitoring templates"
    wb.save(XLSX_OUT)


def main():
    with tempfile.TemporaryDirectory(prefix="fedfina_report_", dir=str(ROOT)) as tmp:
        workdir=Path(tmp)
        crest,letterhead=extract_reference_assets(workdir)
        charts=save_charts(workdir)
        build_pdf(charts,crest,letterhead)
        build_docx(charts,crest,letterhead)
    build_workbook()
    print(f"Created {PDF_OUT.name}")
    print(f"Created {DOCX_OUT.name}")
    print(f"Created {XLSX_OUT.name}")
    print(f"Report pages: {6+len(PAGES)}")


if __name__ == "__main__":
    main()
